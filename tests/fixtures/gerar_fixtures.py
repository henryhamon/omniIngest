"""Gera o corpus de demonstracao do dc.omniIngest, um arquivo por formato suportado.

Reproduzivel de proposito: rodar este script recria o corpus identico, entao a
verificacao de ponta a ponta nao depende de arquivos que alguem arrastou para o
repositorio sem dizer de onde vieram.

    /usr/irissys/bin/irispython tests/fixtures/gerar_fixtures.py

Precisa de python-docx e python-pptx, que ja vem como dependencia do docling.
O PDF e montado byte a byte aqui mesmo: nao ha gerador de PDF na imagem, e
depender de um so para o corpus de teste nao se justifica.
"""

import os

DESTINO = os.path.dirname(os.path.abspath(__file__))

TEXTO_REEMBOLSO = ("O prazo de reembolso e de 30 dias corridos contados a partir da entrega. "
                   "Pedidos fora do prazo sao analisados caso a caso pela equipe de suporte.")
TEXTO_ENTREGA = "O prazo de entrega e de cinco dias uteis para capitais e dez para o interior."
LINHAS_TABELA = [("Categoria", "Prazo", "Taxa"),
                 ("Eletronicos", "30 dias", "0%"),
                 ("Vestuario", "15 dias", "5%"),
                 ("Alimentos", "nao aplicavel", "-")]


def markdown():
    linhas = ["# Manual de Politicas", "", "## Politicas Comerciais", "", "### Reembolso", "",
              TEXTO_REEMBOLSO, "",
              "| " + " | ".join(LINHAS_TABELA[0]) + " |",
              "|" + "|".join(["---"] * 3) + "|"]
    linhas += ["| " + " | ".join(l) + " |" for l in LINHAS_TABELA[1:]]
    linhas += ["", "### Entrega", "", TEXTO_ENTREGA, "", "## Suporte", "",
               "Atendimento de segunda a sexta, das 9h as 18h, pelo canal de chamados."]
    caminho = os.path.join(DESTINO, "manual_politicas.md")
    open(caminho, "w", encoding="utf-8").write("\n".join(linhas) + "\n")
    return caminho


def texto_puro():
    linhas = ["NOTAS DE SUPORTE", "", "Chamado 4412: cliente relatou atraso na entrega da regiao norte.",
              "Chamado 4413: solicitacao de reembolso dentro do prazo de 30 dias, aprovada.",
              "Chamado 4414: duvida sobre taxa de devolucao de vestuario, respondida com a tabela vigente.",
              "", "Formato sem paginas nem cabecalhos: serve para checar que a coluna page_numbers",
              "fica vazia e que o chunking sobrevive a um documento sem estrutura."]
    caminho = os.path.join(DESTINO, "notas_suporte.txt")
    open(caminho, "w", encoding="utf-8").write("\n".join(linhas) + "\n")
    return caminho


def word():
    from docx import Document
    doc = Document()
    doc.add_heading("Contrato de Prestacao de Servicos", level=1)
    doc.add_heading("Clausulas Gerais", level=2)
    doc.add_paragraph("A parte contratante devera observar os procedimentos descritos no anexo tecnico.")
    doc.add_heading("Prazos e Taxas", level=2)
    tabela = doc.add_table(rows=0, cols=3)
    for linha in LINHAS_TABELA:
        celulas = tabela.add_row().cells
        for i, valor in enumerate(linha):
            celulas[i].text = valor
    doc.add_heading("Anexo", level=2)
    doc.add_paragraph("Este anexo trata de assuntos distintos das clausulas gerais acima.")
    caminho = os.path.join(DESTINO, "contrato_servicos.docx")
    doc.save(caminho)
    return caminho


def powerpoint():
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Politicas Comerciais"
    slide.placeholders[1].text = "Resumo para a equipe de atendimento"

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Reembolso e Entrega"
    slide.placeholders[1].text_frame.text = TEXTO_REEMBOLSO
    slide.placeholders[1].text_frame.add_paragraph().text = TEXTO_ENTREGA

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Prazos por Categoria"
    forma = slide.shapes.add_table(len(LINHAS_TABELA), 3, Inches(1), Inches(2), Inches(8), Inches(2))
    for i, linha in enumerate(LINHAS_TABELA):
        for j, valor in enumerate(linha):
            forma.table.cell(i, j).text = valor
    caminho = os.path.join(DESTINO, "apresentacao_politicas.pptx")
    prs.save(caminho)
    return caminho


def pdf():
    """PDF de uma pagina com texto e uma tabela desenhada com linhas de verdade.

    Montado na mao porque nao ha biblioteca de escrita de PDF na imagem. O
    conteudo e texto vetorial (nao imagem), que e o caso que importa: o pipeline
    roda com OCR desligado de proposito.
    """
    def escapar(t):
        return t.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")

    partes = ["BT /F1 18 Tf 60 780 Td (Tabela de Precos e Prazos) Tj ET",
              "BT /F1 11 Tf 60 750 Td (" + escapar(TEXTO_ENTREGA) + ") Tj ET"]

    topo, altura, larguras, x0 = 700, 24, [160, 150, 100], 60
    total = sum(larguras)
    # grade da tabela
    for i in range(len(LINHAS_TABELA) + 1):
        y = topo - i * altura
        partes.append(f"{x0} {y} m {x0 + total} {y} l S")
    x = x0
    for largura in larguras + [0]:
        partes.append(f"{x} {topo} m {x} {topo - len(LINHAS_TABELA) * altura} l S")
        x += largura
    # texto das celulas
    for i, linha in enumerate(LINHAS_TABELA):
        y = topo - (i + 1) * altura + 8
        x = x0 + 5
        for j, valor in enumerate(linha):
            partes.append(f"BT /F1 10 Tf {x} {y} Td ({escapar(valor)}) Tj ET")
            x += larguras[j]
    partes.append("BT /F1 11 Tf 60 560 Td (" + escapar(TEXTO_REEMBOLSO[:90]) + ") Tj ET")
    conteudo = "0.6 w\n" + "\n".join(partes)

    objetos = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        "/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(conteudo)} >>\nstream\n{conteudo}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]

    saida = bytearray(b"%PDF-1.4\n")
    deslocamentos = []
    for numero, corpo in enumerate(objetos, start=1):
        deslocamentos.append(len(saida))
        saida += f"{numero} 0 obj\n{corpo}\nendobj\n".encode("latin-1")
    inicio_xref = len(saida)
    saida += f"xref\n0 {len(objetos) + 1}\n".encode("latin-1")
    saida += b"0000000000 65535 f \n"
    for deslocamento in deslocamentos:
        saida += f"{deslocamento:010d} 00000 n \n".encode("latin-1")
    saida += (f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\n"
              f"startxref\n{inicio_xref}\n%%EOF\n").encode("latin-1")

    caminho = os.path.join(DESTINO, "tabela_precos.pdf")
    open(caminho, "wb").write(bytes(saida))
    return caminho


if __name__ == "__main__":
    for gerar in (markdown, texto_puro, word, powerpoint, pdf):
        caminho = gerar()
        print("%-40s %7d bytes" % (os.path.basename(caminho), os.path.getsize(caminho)))
